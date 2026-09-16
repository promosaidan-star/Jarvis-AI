"""Blind human-rating packet (metric 4) + scorer.

build:  python eval/rating_packet.py build run_2026-09-11
  -> results/<run>/rating/rating_sheet.xlsx (one tab per rater, same 15 cards, shuffled per rater)
     results/<run>/rating/KEY_do_not_open.json (card id -> symbol, draft|verified)
score:  python eval/rating_packet.py score run_2026-09-11
  -> reads the filled sheet, prints clarity + trust by version and inter-rater agreement

Cards come only from the 20 positions the verifier corrected (identical draft/verified
pairs would say nothing). 15 symbols; alternating draft/verified (8 drafts, 7 verified).
Every rater rates all 15 so agreement can be computed on the same items.
"""
from __future__ import annotations

import itertools
import json
import random
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
from run_eval import load  # noqa: E402

RATERS = ["Rater A", "Rater B", "Rater C"]


def build(run: Path, out: Path):
    glossary, rows = load(run)
    rng = random.Random(917)
    corrected = [r for r in rows if r["ok"] is False]
    chosen = rng.sample(corrected, 15)
    cards = []
    for i, r in enumerate(chosen):
        version = "draft" if i % 2 == 0 else "verified"
        c = r[version]
        p = r["pos"]
        cards.append({"id": f"C{i + 1:02d}", "symbol": p["symbol"], "version": version,
                      "headline": c["headline"], "body": c["body"]})
    rng.shuffle(cards)
    ids = {c["id"]: c for c in cards}
    out.mkdir(parents=True, exist_ok=True)
    (out / "KEY_do_not_open.json").write_text(json.dumps({k: {"symbol": v["symbol"], "version": v["version"]} for k, v in ids.items()}, indent=1), encoding="utf-8")

    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    for line in [
        "Blind rating: Book Explainer cards (about 25 minutes)",
        "",
        "Each card explains one position in a paper-trading book. Some cards were checked and corrected, some were not; you are not told which.",
        "Rate each card on your own tab only. Do not discuss cards with the other raters until everyone is done.",
        "Clarity (1-5): 1 = I could not follow it, 3 = understandable with effort, 5 = clear on first read.",
        "Trust (Yes/No): would you paste this text into a risk memo without re-checking it yourself?",
        "Optional comment: anything that looked wrong, vague or over-claimed.",
    ]:
        ins.append([line])
    ins["A1"].font = Font(bold=True, size=13)
    ins.column_dimensions["A"].width = 140

    for rater in RATERS:
        ws = wb.create_sheet(rater)
        ws.append(["Card", "Headline", "Explanation", "Clarity (1-5)", "Trust in risk memo (Yes/No)", "Comment"])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3A5F")
        order = cards[:]
        random.Random(rater).shuffle(order)
        for c in order:
            ws.append([c["id"], c["headline"], c["body"], None, None, None])
        for col, w in zip("ABCDEF", (7, 40, 110, 12, 16, 40)):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        dv1 = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
        dv2 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
        ws.add_data_validation(dv1)
        ws.add_data_validation(dv2)
        dv1.add(f"D2:D{len(order) + 1}")
        dv2.add(f"E2:E{len(order) + 1}")
        ws.freeze_panes = "B2"
    wb.save(out / "rating_sheet.xlsx")
    print("wrote", out / "rating_sheet.xlsx", "and KEY_do_not_open.json")


def score(out: Path):
    key = json.loads((out / "KEY_do_not_open.json").read_text(encoding="utf-8"))
    wb = load_workbook(out / "rating_sheet.xlsx", data_only=True)
    ratings = {}
    for rater in RATERS:
        for cid, _, _, clar, trust, _ in wb[rater].iter_rows(min_row=2, values_only=True):
            if cid and clar is not None and trust:
                ratings.setdefault(cid, {})[rater] = (int(clar), str(trust).strip().lower().startswith("y"))
    res = {}
    for version in ("draft", "verified"):
        cl = [v[0] for cid, rs in ratings.items() if key[cid]["version"] == version for v in rs.values()]
        tr = [v[1] for cid, rs in ratings.items() if key[cid]["version"] == version for v in rs.values()]
        res[version] = {"n_ratings": len(cl), "mean_clarity": round(sum(cl) / len(cl), 2) if cl else None,
                        "trust_yes_rate": round(sum(tr) / len(tr), 2) if tr else None}
    # pairwise agreement on trust (share of rater pairs agreeing) + within-1-point clarity agreement
    agree_t, agree_c, pairs = 0, 0, 0
    for rs in ratings.values():
        for a, b in itertools.combinations(rs.values(), 2):
            pairs += 1
            agree_t += a[1] == b[1]
            agree_c += abs(a[0] - b[0]) <= 1
    res["agreement"] = {"rater_pairs": pairs, "trust_same_answer": round(agree_t / pairs, 2) if pairs else None,
                        "clarity_within_1_point": round(agree_c / pairs, 2) if pairs else None}
    (out / "rating_results.json").write_text(json.dumps(res, indent=2), encoding="utf-8")
    print(json.dumps(res, indent=2))


if __name__ == "__main__":
    cmd, run = sys.argv[1], Path(sys.argv[2])
    out = Path("results") / run.name / "rating"
    build(run, out) if cmd == "build" else score(out)
