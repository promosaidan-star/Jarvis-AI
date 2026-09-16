"""Make a PowerPoint deck from an outline. Two ways in:

    python -m jarvis.pptx_builder outline.json out/deck.pptx
    python -m jarvis.pptx_builder --topic "Book Explainer: results" --slides 6 out/deck.pptx

Outline JSON:
{"title": "...", "subtitle": "...",
 "slides": [{"title": "...", "bullets": ["...", "..."], "notes": "speaker notes (optional)"},
            {"title": "...", "table": {"header": [...], "rows": [[...], ...]}},
            {"title": "...", "image": "path/to.png", "caption": "..."}]}

With --topic, the outline is drafted by the LLM first (and saved next to the deck so it
can be edited and re-run). The deck is 16:9, one typeface, dark title bar, no clip art.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

INK = RGBColor(0x12, 0x20, 0x2C)
ACCENT = RGBColor(0x0F, 0x6A, 0x8B)
MUTED = RGBColor(0x5A, 0x6B, 0x7B)
FONT = "Calibri"


def _text(frame, text, size, bold=False, color=INK):
    frame.clear()
    p = frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(size), bold, FONT, color
    return p


def _bar(slide, prs, title):
    bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = INK
    bar.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), prs.slide_width - Inches(1), Inches(0.7))
    _text(tb.text_frame, title, 26, True, RGBColor(0xFF, 0xFF, 0xFF))


def build(outline: dict, out: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    band = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    band.fill.solid()
    band.fill.fore_color.rgb = INK
    band.line.fill.background()
    tb = s.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.5), Inches(1.5))
    _text(tb.text_frame, outline.get("title", "Untitled"), 40, True, RGBColor(0xFF, 0xFF, 0xFF))
    tb.text_frame.word_wrap = True
    if outline.get("subtitle"):
        sb = s.shapes.add_textbox(Inches(0.8), Inches(4.1), Inches(11.5), Inches(1))
        _text(sb.text_frame, outline["subtitle"], 18, False, RGBColor(0xC8, 0xD4, 0xDE))

    for spec in outline.get("slides", []):
        s = prs.slides.add_slide(blank)
        _bar(s, prs, spec.get("title", ""))
        top, left, width = Inches(1.5), Inches(0.7), prs.slide_width - Inches(1.4)
        if spec.get("bullets"):
            tb = s.shapes.add_textbox(left, top, width, Inches(5.3))
            tf = tb.text_frame
            tf.word_wrap = True
            for i, b in enumerate(spec["bullets"]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                r = p.add_run()
                r.text = "•  " + b
                r.font.size, r.font.name, r.font.color.rgb = Pt(20), FONT, INK
                p.space_after = Pt(10)
        if spec.get("table"):
            t = spec["table"]
            rows, cols = len(t["rows"]) + 1, len(t["header"])
            shape = s.shapes.add_table(rows, cols, left, top, width, Emu(0))
            table = shape.table
            for j, h in enumerate(t["header"]):
                c = table.cell(0, j)
                c.text = str(h)
                c.fill.solid()
                c.fill.fore_color.rgb = ACCENT
                for p in c.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.bold, r.font.size, r.font.color.rgb = True, Pt(14), RGBColor(0xFF, 0xFF, 0xFF)
            for i, row in enumerate(t["rows"], start=1):
                for j, v in enumerate(row):
                    c = table.cell(i, j)
                    c.text = str(v)
                    for p in c.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.size, r.font.name = Pt(13), FONT
        if spec.get("image") and Path(spec["image"]).exists():
            pic = s.shapes.add_picture(spec["image"], left, top, height=Inches(5.0))
            if pic.width > width:
                pic.width, pic.height = width, int(pic.height * width / pic.width)
            if spec.get("caption"):
                cb = s.shapes.add_textbox(left, Inches(6.7), width, Inches(0.5))
                _text(cb.text_frame, spec["caption"], 12, False, MUTED)
        if spec.get("notes"):
            s.notes_slide.notes_text_frame.text = spec["notes"]
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


def draft_outline(topic: str, n_slides: int, context: str = "") -> dict:
    from .llm import ask_json
    system = ("Draft a slide outline as JSON {title, subtitle, slides:[{title, bullets:[3-5 short bullets], notes}]}. "
              f"Exactly {n_slides} content slides. Concrete, no filler bullets like 'Introduction'. Plain language; "
              "numbers before adjectives.")
    return ask_json(system, f"Topic: {topic}\n\nContext:\n{context[:6000]}")


if __name__ == "__main__":
    args = sys.argv[1:]
    if "--topic" in args:
        topic = args[args.index("--topic") + 1]
        n = int(args[args.index("--slides") + 1]) if "--slides" in args else 6
        out = Path(args[-1])
        ctx = ""
        if "--context" in args:
            ctx = Path(args[args.index("--context") + 1]).read_text(encoding="utf-8")
        outline = draft_outline(topic, n, ctx)
        out.with_suffix(".outline.json").write_text(json.dumps(outline, indent=1), encoding="utf-8")
    else:
        outline, out = json.loads(Path(args[0]).read_text(encoding="utf-8")), Path(args[1])
    print(build(outline, out))
